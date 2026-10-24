"""第四阶段回归测试：PPTX、扫描 PDF/OCR、跨文档检索问答与版本失效。"""

import io
import unittest
from unittest.mock import patch

import pytesseract
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from app import converters, migrate, store

from test_phase2 import BaseTest, make_docx, make_pdf


def make_png(size=(90, 60), color=(20, 120, 140)) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="PNG")
    return buf.getvalue()


def make_pptx() -> bytes:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "注意力机制"
    slide.placeholders[1].text = "第一段说明\n- 要点一\n- 要点二"
    slide.shapes.add_picture(
        io.BytesIO(make_png()), Inches(1), Inches(4), width=Inches(1)
    )
    table = slide.shapes.add_table(
        2, 2, Inches(4), Inches(4), Inches(3), Inches(1)
    ).table
    table.cell(0, 0).text = "列一"
    table.cell(0, 1).text = "列二"
    table.cell(1, 0).text = "值一"
    table.cell(1, 1).text = "值二"
    slide.notes_slide.notes_text_frame.text = "讲课时强调权重归一化"

    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "小结"
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def make_scanned_pdf() -> bytes:
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf)
    pdf.drawImage(ImageReader(io.BytesIO(make_png((400, 120)))), 72, 600, width=300, height=90)
    pdf.showPage()
    pdf.save()
    return buf.getvalue()


class PptxTests(BaseTest):
    def test_pptx_imports_text_table_image_and_notes(self):
        response = self.upload("deck.pptx", make_pptx())
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        doc_id = response.get_json()["document_id"]
        content = self.doc_content(doc_id)
        self.assertEqual(content["status"], "ready")
        self.assertEqual(content["view_mode"], "pdf")
        self.assertEqual(len(content["chapters"]), 2)
        self.assertIn("<table>", content["html"])
        self.assertIn("<img", content["html"])
        self.assertIn("权重归一化", content["html"])

        image_src = None
        for block in content["blocks"]:
            if block.get("type") == "image":
                image_src = block.get("id")
        self.assertTrue(image_src)
        asset = self.client.get(f"/api/documents/{doc_id}/assets/slide1_3.png")
        self.assertEqual(asset.status_code, 200)
        self.assertIn(b"PNG", asset.data[:8] or b"")
        asset.close()

        labels = {block.get("locator_label") for block in content["blocks"]}
        self.assertTrue(any("第 1 页" in (label or "") for label in labels))

    def test_docx_conversion_failure_is_reported(self):
        with patch(
            "app.converters.convert_to_pdf",
            side_effect=converters.ConvertError("需要本机 Microsoft Office"),
        ):
            self.upload("docx.docx", make_docx())
        doc = self.docs()[0]
        self.assertEqual(doc["version_status"], "failed")
        self.assertIn("Office", doc.get("version_error") or "")

    def test_pptx_missing_dependency_is_reported(self):
        with patch.dict("sys.modules", {"pptx": None}):
            response = self.upload("broken.pptx", b"not a deck")
        self.assertEqual(response.status_code, 200)
        doc = self.docs()[0]
        self.assertEqual(doc["version_status"], "failed")
        self.assertIn("PPTX", doc.get("version_error") or "")


class ScanPdfTests(BaseTest):
    def test_scanned_pdf_ocr_with_mocked_engine(self):
        with patch("pytesseract.image_to_string", return_value="扫描页识别出的中文内容\n第二段"):
            response = self.upload("scan.pdf", make_scanned_pdf())
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        doc_id = response.get_json()["document_id"]
        content = self.doc_content(doc_id)
        self.assertEqual(content["status"], "ready")
        self.assertIn("扫描页识别出的中文内容", content["html"])
        locators = [block.get("locator") for block in content["blocks"]]
        self.assertTrue(any((locator or {}).get("ocr") for locator in locators))

    def test_missing_tesseract_reports_clear_error(self):
        with patch(
            "pytesseract.image_to_string",
            side_effect=pytesseract.TesseractNotFoundError(),
        ):
            response = self.upload("scan.pdf", make_scanned_pdf())
        self.assertEqual(response.status_code, 200)
        doc = self.docs()[0]
        self.assertEqual(doc["version_status"], "failed")
        message = doc.get("version_error") or ""
        self.assertIn("Tesseract", message)
        self.assertIn("扫描", message)


class RenderMigrationTests(BaseTest):
    def test_existing_office_documents_get_playback_pdf(self):
        doc_id = self.upload("old.docx", make_docx()).get_json()["document_id"]
        content = self.doc_content(doc_id)
        version_id = content["version"]["id"]
        store.set_version_render(version_id, None)
        missing = self.doc_content(doc_id)
        self.assertEqual(missing["view_mode"], "markdown")
        self.assertTrue(missing["pdf_missing"])

        result = migrate.ensure_render_paths()
        self.assertGreaterEqual(result["updated"], 1)
        restored = self.doc_content(doc_id)
        self.assertEqual(restored["view_mode"], "pdf")
        self.assertFalse(restored["pdf_missing"])

    def test_existing_pdf_gets_playback_path(self):
        doc_id = self.upload("old.pdf", make_pdf()).get_json()["document_id"]
        version_id = self.doc_content(doc_id)["version"]["id"]
        store.set_version_render(version_id, None)
        migrate.ensure_render_paths()
        content = self.doc_content(doc_id)
        self.assertEqual(content["view_mode"], "pdf")
        self.assertTrue(content["pdf_url"])


class SearchTests(BaseTest):
    def setUp(self):
        super().setUp()
        (self.library / "alpha.md").write_text(
            "# 注意力\n\n注意力机制按相关性加权求和。\n\n## 训练\n\n训练需要梯度。",
            encoding="utf-8",
        )
        (self.library / "beta.md").write_text(
            "# 数据库\n\n数据库用索引加速查询。", encoding="utf-8"
        )
        self.client.post("/api/library/scan")
        by_name = {doc["source_name"]: doc["id"] for doc in self.docs()}
        self.alpha = by_name["alpha.md"]
        self.beta = by_name["beta.md"]

    def test_search_returns_relevant_document_and_locator(self):
        data = self.client.get("/api/search", query_string={"q": "注意力"}).get_json()
        self.assertGreaterEqual(data["count"], 1)
        self.assertTrue(all(item["document_id"] == self.alpha for item in data["results"]))
        result = data["results"][0]
        self.assertIn("注意力", result["snippet"])
        self.assertIn("行", result["locator"])

    def test_search_scope_and_empty_query(self):
        scoped = self.client.get(
            "/api/search", query_string={"q": "注意力", "doc_id": self.beta}
        ).get_json()
        self.assertEqual(scoped["count"], 0)
        empty = self.client.get("/api/search", query_string={"q": "  "}).get_json()
        self.assertEqual(empty["count"], 0)

    def test_search_ignores_old_versions(self):
        path = self.library / "alpha.md"
        path.write_text("# 新内容\n\n换成支持向量机。", encoding="utf-8")
        self.client.post("/api/library/scan")
        old = self.client.get("/api/search", query_string={"q": "梯度"}).get_json()
        self.assertEqual(old["count"], 0)
        new = self.client.get("/api/search", query_string={"q": "支持向量机"}).get_json()
        self.assertEqual(new["results"][0]["document_id"], self.alpha)


class CrossDocumentAskTests(BaseTest):
    def setUp(self):
        super().setUp()
        self.alpha = self.upload(
            "alpha.md", "# 模型\n\n注意力模型按权重聚合。".encode("utf-8")
        ).get_json()["document_id"]
        self.beta = self.upload(
            "beta.md", "# 系统\n\n数据库模型负责存储。".encode("utf-8")
        ).get_json()["document_id"]
        self.client.post(
            "/api/config",
            json={
                "api_base": "https://api.example.com",
                "api_key": "sk-phase4-example",
                "model": "demo",
                "timeout": 20,
                "temperature": 0,
            },
        )

    def test_ask_cites_multiple_documents(self):
        found = self.client.get("/api/search", query_string={"q": "模型"}).get_json()
        ids = {item["document_id"]: item["block_id"] for item in found["results"]}
        self.assertEqual(set(ids), {self.alpha, self.beta})
        reply = f"两份资料都相关 [cite:{ids[self.alpha]}] [cite:{ids[self.beta]}]"
        with patch("app.llm.chat", return_value=reply):
            response = self.client.post("/api/ask", json={"query": "模型"})
        self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
        data = response.get_json()
        self.assertEqual(len(data["citations"]), 2)
        self.assertEqual({source["document_id"] for source in data["sources"]}, {self.alpha, self.beta})

    def test_ask_drops_forged_citation(self):
        with patch("app.llm.chat", return_value=f"回答 [cite:{'f' * 16}]"):
            response = self.client.post("/api/ask", json={"query": "模型"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.get_json()["citations"], [])
        self.assertNotIn("ffffffff", response.get_json()["answer"])

    def test_ask_requires_query(self):
        response = self.client.post("/api/ask", json={"query": " "})
        self.assertEqual(response.status_code, 400)


class VersionInvalidationTests(BaseTest):
    def test_new_version_clears_prereq_and_keeps_old_learning_records(self):
        path = self.library / "v.md"
        path.write_text("# 第一版\n\n旧内容梯度。", encoding="utf-8")
        self.client.post("/api/library/scan")
        doc = self.docs()[0]
        doc_id = doc["id"]
        first = self.doc_content(doc_id)
        first_version = first["version"]["id"]
        block_id = first["blocks"][0]["id"]
        store.save_prereq(doc_id, {"focus": "基础", "questions": []})
        store.save_understanding(
            doc_id, first_version, block_id, "", "问题", "回答", "反馈", "理解"
        )

        path.write_text("# 第二版\n\n新内容支持向量机。", encoding="utf-8")
        self.client.post("/api/library/scan")
        content = self.doc_content(doc_id)
        self.assertEqual(content["version"]["number"], 2)
        self.assertEqual(len(content["versions"]), 2)
        self.assertIsNone(store.get_prereq(doc_id))

        learning = self.client.get(f"/api/documents/{doc_id}/learning").get_json()
        self.assertEqual(len(learning["understanding"]), 1)
        self.assertEqual(learning["understanding"][0]["version_id"], first_version)
        self.assertNotEqual(content["version"]["id"], first_version)


class FrontendPhase4Tests(unittest.TestCase):
    def test_phase4_dom_ids_exist(self):
        from pathlib import Path

        from app import config

        html = (Path(config.APP_DIR) / "app" / "templates" / "index.html").read_text(
            encoding="utf-8"
        )
        for element_id in (
            "search-panel",
            "search-results",
            "btn-ask-cross",
            "doc-version-note",
            "pdf-view",
            "pdf-frame",
        ):
            self.assertIn(f'id="{element_id}"', html, element_id)
        self.assertIn(".pptx", html)


if __name__ == "__main__":
    unittest.main()
