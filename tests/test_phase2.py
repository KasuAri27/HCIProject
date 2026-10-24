"""第二阶段回归测试：导入、解析、版本、迁移、位置与标注。"""

import io
import json
import os
import re
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from reportlab.pdfgen import canvas

from app import (
    api_config,
    config,
    imports,
    migrate,
    parsers,
    source,
    store,
    web,
)


def make_docx(text: str = "DOCX 第一段") -> bytes:
    doc = Document()
    doc.add_heading("DOCX 标题", level=1)
    doc.add_paragraph(text)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "列一"
    table.cell(0, 1).text = "列二"
    table.cell(1, 0).text = "值一"
    table.cell(1, 1).text = "值二"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_pdf(text: str = "PDF first paragraph.") -> bytes:
    buf = io.BytesIO()
    pdf = canvas.Canvas(buf)
    pdf.drawString(72, 720, text)
    pdf.drawString(72, 700, "PDF second line.")
    pdf.showPage()
    pdf.drawString(72, 720, "PDF page two.")
    pdf.save()
    return buf.getvalue()


class BaseTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)
        self.library = self.tmp / "library"
        self.library.mkdir()
        (self.tmp / "db").mkdir()
        self._env = patch.dict(os.environ, {"LLM_API_KEY": ""}, clear=False)
        self._env.start()
        self._patchers = [
            patch.object(store, "DB_PATH", self.tmp / "db" / "learning.db"),
            patch.object(source, "LIBRARY_DIR", self.library),
            patch.object(config, "APP_DIR", self.tmp),
            patch.object(config, "ORIGINALS_DIR", self.tmp / "db" / "originals"),
            patch.object(config, "BACKUP_DIR", self.tmp / "db" / "backups"),
            patch.object(api_config, "CONFIG_PATH", self.tmp / "db" / "api_config.json"),
            patch("app.imports.start_job", side_effect=self._run_job),
            patch("app.converters.convert_to_pdf", side_effect=self._fake_convert),
        ]
        for p in self._patchers:
            p.start()
        web._db_ready_path = ""
        store.init_db()
        web.app.config.update(TESTING=True)
        self.client = web.app.test_client()

    def _run_job(self, job_id):
        imports.process_job(job_id)

    def _fake_convert(self, source, target):
        target = Path(target)
        target.parent.mkdir(parents=True, exist_ok=True)
        pdf = canvas.Canvas(str(target))
        pdf.drawString(72, 720, "converted")
        pdf.showPage()
        pdf.save()
        return target

    def tearDown(self):
        for p in reversed(self._patchers):
            p.stop()
        self._env.stop()
        self._tmp.cleanup()

    def upload(self, name: str, data: bytes):
        return self.client.post(
            "/api/import",
            data={"file": (io.BytesIO(data), name)},
            content_type="multipart/form-data",
        )

    def docs(self):
        return self.client.get("/api/library").get_json()["documents"]

    def doc_content(self, doc_id):
        return self.client.get(f"/api/documents/{doc_id}").get_json()


class ParserTests(unittest.TestCase):
    def test_markdown_structure_and_locators(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "a.md"
            path.write_text(
                "# 第一章\n\n段落一\n\n## 小节\n\n- a\n- b\n\n```py\nprint(1)\n```\n\n| A | B |\n|---|---|\n| 1 | 2 |\n",
                encoding="utf-8",
            )
            parsed = parsers.parse_document(path, "md", "a")
            self.assertGreaterEqual(len(parsed["chapters"]), 2)
            types = [b["type"] for c in parsed["chapters"] for b in c["blocks"]]
            self.assertIn("heading", types)
            self.assertIn("list", types)
            self.assertIn("code", types)
            self.assertIn("table", types)
            locator = parsed["chapters"][0]["blocks"][0]["source_locator"]
            self.assertEqual(locator["line_start"], 1)
            self.assertIn("第 1 行", parsers.locator_label("md", locator))
            self.assertTrue(parsed["markdown"].strip())
            self.assertTrue(all(b.get("markdown") for c in parsed["chapters"] for b in c["blocks"]))


class ImportTests(BaseTest):
    def test_imports_markdown_txt_docx_pdf(self):
        cases = [
            ("a.md", "# MD\n\nHello **world**.".encode("utf-8"), "md"),
            ("b.txt", "纯文本\n第二行".encode("utf-8"), "txt"),
            ("c.docx", make_docx(), "docx"),
            ("d.pdf", make_pdf(), "pdf"),
        ]
        for name, data, fmt in cases:
            response = self.upload(name, data)
            self.assertEqual(response.status_code, 200, response.get_data(as_text=True))
            payload = response.get_json()
            self.assertFalse(payload["duplicate"])
            doc = self.client.get(f"/api/documents/{payload['document_id']}").get_json()
            self.assertEqual(doc["status"], "ready")
            self.assertTrue(doc["html"])
            self.assertEqual(doc["version"]["source_format"], fmt)
            self.assertTrue(doc["blocks"])
            expected_view = "markdown" if fmt in ("md", "txt") else "pdf"
            self.assertEqual(doc["view_mode"], expected_view, fmt)
            self.assertTrue(doc["pdf_url"] if expected_view == "pdf" else not doc["pdf_url"])
            if expected_view == "pdf":
                playable = self.client.get(doc["pdf_url"])
                self.assertEqual(playable.status_code, 200)
                self.assertEqual(playable.mimetype, "application/pdf")
                self.assertIn(b"%PDF", playable.data[:5])
                playable.close()

        documents = self.docs()
        self.assertEqual(len(documents), 4)
        self.assertEqual({d["source_format"] for d in documents}, {"md", "txt", "docx", "pdf"})
        for doc in documents:
            original = self.client.get(f"/api/documents/{doc['id']}/original")
            self.assertEqual(original.status_code, 200)
            original.close()

    def test_plain_text_is_previewed_not_shown_as_source(self):
        response = self.upload("plain.txt", b"*not italic* and # not heading")
        doc_id = response.get_json()["document_id"]
        html = self.doc_content(doc_id)["html"]
        self.assertIn("*not italic*", html)
        self.assertNotIn("<em>", html)
        self.assertNotIn('class="plain-text"', html)

    def test_markdown_preview_and_docx_table(self):
        md_id = self.upload("preview.md", b"# Title\n\n**bold**").get_json()["document_id"]
        md_html = self.doc_content(md_id)["html"]
        self.assertIn("<h1", md_html)
        self.assertIn("<strong>bold</strong>", md_html)

        docx_id = self.upload("table.docx", make_docx()).get_json()["document_id"]
        docx_html = self.doc_content(docx_id)["html"]
        self.assertIn("<h1", docx_html)
        self.assertIn("<table>", docx_html)

    def test_duplicate_import_is_idempotent(self):
        first = self.upload("a.md", b"# A\n\nsame").get_json()
        second = self.upload("a.md", b"# A\n\nsame").get_json()
        self.assertTrue(second["duplicate"])
        self.assertEqual(first["document_id"], second["document_id"])
        self.assertEqual(len(self.docs()), 1)

    def test_version_and_rename_keep_document_id_and_notes(self):
        path = self.library / "old.md"
        path.write_text("# Old\n\nfirst", encoding="utf-8")
        self.client.post("/api/library/scan")
        doc = self.docs()[0]
        doc_id = doc["id"]
        self.client.post(f"/api/documents/{doc_id}/open")
        self.client.put("/api/notes", json={"doc_id": doc_id, "content": "keep me"})

        path.rename(self.library / "renamed.md")
        self.client.post("/api/library/scan")
        renamed = self.docs()[0]
        self.assertEqual(renamed["id"], doc_id)
        self.assertEqual(renamed["source_path"], str((self.library / "renamed.md").resolve()))
        self.assertEqual(self.client.get(f"/api/notes?doc_id={doc_id}").get_json()["content"], "keep me")

        (self.library / "renamed.md").write_text("# New\n\nsecond", encoding="utf-8")
        self.client.post("/api/library/scan")
        updated = self.docs()[0]
        self.assertEqual(updated["id"], doc_id)
        self.assertEqual(updated["version_count"], 2)
        self.assertEqual(updated["version_no"], 2)
        self.assertIn("second", self.doc_content(doc_id)["html"])

    def test_failed_parse_can_retry(self):
        path = self.library / "broken.md"
        path.write_text("# Broken", encoding="utf-8")
        with patch("app.parsers.parse_document", side_effect=parsers.ParseError("坏文件")):
            self.client.post("/api/library/scan")
        doc = self.docs()[0]
        self.assertEqual(doc["version_status"], "failed")
        job = doc.get("job") or {}
        self.assertTrue(job.get("id"))

        self.client.post(f"/api/imports/{job['id']}/retry")
        doc = self.docs()[0]
        self.assertEqual(doc["version_status"], "ready")
        self.assertEqual(self.doc_content(doc["id"])["status"], "ready")

    def test_scanned_pdf_returns_clear_error(self):
        buf = io.BytesIO()
        pdf = canvas.Canvas(buf)
        pdf.showPage()
        pdf.save()
        response = self.upload("scan.pdf", buf.getvalue())
        self.assertEqual(response.status_code, 200)
        doc_id = response.get_json()["document_id"]
        content = self.client.get(f"/api/documents/{doc_id}").get_json()
        self.assertEqual(content["status"], "failed")
        self.assertIn("PDF", content["error"])


class ReadingTests(BaseTest):
    def setUp(self):
        super().setUp()
        response = self.upload("read.md", "# Read\n\n第一段内容。\n\n第二段内容。".encode("utf-8"))
        self.doc_id = response.get_json()["document_id"]

    def test_no_api_reading_and_notes_work(self):
        session = self.client.get("/api/session").get_json()
        self.assertFalse(session["llm_configured"])
        opened = self.client.post(f"/api/documents/{self.doc_id}/open")
        self.assertEqual(opened.status_code, 200)
        content = self.doc_content(self.doc_id)
        self.assertEqual(content["status"], "ready")
        note = self.client.put(
            "/api/notes", json={"doc_id": self.doc_id, "content": "离线笔记"}
        )
        self.assertEqual(note.status_code, 200)
        self.assertEqual(
            self.client.get(f"/api/notes?doc_id={self.doc_id}").get_json()["content"],
            "离线笔记",
        )

    def test_position_is_persisted(self):
        content = self.doc_content(self.doc_id)
        block = content["blocks"][1]
        position = {
            "version_id": content["version"]["id"],
            "block_id": block["id"],
            "block_offset": 0.5,
            "scroll_ratio": 0.25,
        }
        saved = self.client.post(f"/api/documents/{self.doc_id}/position", json=position)
        self.assertEqual(saved.status_code, 200)
        loaded = self.client.get(f"/api/documents/{self.doc_id}").get_json()["position"]
        self.assertEqual(loaded["block_id"], block["id"])
        self.assertAlmostEqual(loaded["block_offset"], 0.5)

    def test_annotations_crud_and_locator(self):
        content = self.doc_content(self.doc_id)
        block = content["blocks"][1]
        created = self.client.post(
            f"/api/documents/{self.doc_id}/annotations",
            json={
                "doc_id": self.doc_id,
                "block_id": block["id"],
                "start_offset": 0,
                "end_offset": 3,
                "quote": "第一段",
                "note": "批注",
                "color": "amber",
            },
        )
        self.assertEqual(created.status_code, 200)
        annotation = created.get_json()["annotation"]
        self.assertIn("第", content["blocks"][1]["locator_label"])

        listed = self.client.get(f"/api/documents/{self.doc_id}").get_json()["annotations"]
        self.assertEqual(len(listed), 1)
        self.assertTrue(listed[0]["locator_label"])

        patched = self.client.patch(
            f"/api/annotations/{annotation['id']}",
            json={"note": "改过", "color": "teal"},
        )
        self.assertEqual(patched.status_code, 200)
        self.assertEqual(patched.get_json()["annotation"]["note"], "改过")

        deleted = self.client.delete(f"/api/annotations/{annotation['id']}")
        self.assertEqual(deleted.status_code, 200)
        self.assertEqual(self.client.get(f"/api/documents/{self.doc_id}").get_json()["annotations"], [])

    def test_original_file_is_served(self):
        response = self.client.get(f"/api/documents/{self.doc_id}/original")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Read", response.data)
        response.close()


class MigrationTests(BaseTest):
    def test_migrates_json_with_backup_and_is_idempotent(self):
        (self.library / "old.md").write_text("# Old\n\ncontent", encoding="utf-8")
        data_dir = self.tmp / "db"
        (data_dir / "progress.json").write_text(
            json.dumps(
                {
                    "selected_source": {"rel": "old.md", "title": "Old"},
                    "read_rels": {"old.md": {"title": "Old", "read_at": "x"}},
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        (data_dir / "chapter_notes.json").write_text(
            json.dumps({"old.md": {"title": "Old", "content": "migrated note"}}, ensure_ascii=False),
            encoding="utf-8",
        )
        (data_dir / "chapter_chats.json").write_text(
            json.dumps(
                {
                    "old.md": {
                        "messages": [
                            {"role": "user", "content": "hi", "via": "chat"},
                            {"role": "assistant", "content": "hello", "via": "chat"},
                        ]
                    }
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        result = migrate.migrate_if_needed()
        self.assertTrue(result["migrated"])
        self.assertTrue(Path(result["backup"]).is_dir())
        self.assertTrue((Path(result["backup"]) / "progress.json").exists())
        doc_id = self.docs()[0]["id"]
        self.assertEqual(store.get_selected_document_id(), doc_id)
        self.assertEqual(store.get_note(doc_id)["content"], "migrated note")
        self.assertEqual(len(store.get_chat(doc_id)), 2)
        second = migrate.migrate_if_needed()
        self.assertFalse(second["migrated"])

    def test_backfills_markdown_for_existing_versions(self):
        doc_id = self.upload("legacy.txt", b"*literal* text").get_json()["document_id"]
        conn = store.connect()
        conn.execute("UPDATE document_versions SET markdown = NULL")
        conn.execute("UPDATE blocks SET markdown = NULL")
        conn.commit()
        conn.close()
        count = migrate.normalize_markdown()
        self.assertEqual(count, 1)
        content = self.doc_content(doc_id)
        self.assertTrue(all(b.get("id") for b in content["blocks"]))
        self.assertIn("<p>*literal* text</p>", content["html"])


class FrontendContractTests(unittest.TestCase):
    def test_dom_ids_exist(self):
        root = Path(__file__).resolve().parent.parent
        js = (root / "app/static/app.js").read_text(encoding="utf-8")
        html = (root / "app/templates/index.html").read_text(encoding="utf-8")
        ids = set(re.findall(r'id="([^"]+)"', html))
        refs = set(re.findall(r'\$\("([^"]+)"\)', js))
        dynamic = {"btn-empty-import", "btn-empty-scan", "btn-status-retry"}
        self.assertEqual(refs - ids - dynamic, set())
        self.assertIn('id="doc-original"', html)
        self.assertIn("下载原文件", html)
        self.assertIn("download", html)


if __name__ == "__main__":
    unittest.main()


class AIAndConfigTests(BaseTest):
    def setUp(self):
        super().setUp()
        self.doc_id = self.upload("ai.md", "# AI\n\n资料正文。".encode("utf-8")).get_json()["document_id"]
        response = self.client.post(
            "/api/config",
            json={
                "api_base": "https://api.example.com",
                "api_key": "sk-abcdefghijklmnop",
                "model": "demo",
                "timeout": 20,
                "temperature": 0,
            },
        )
        self.assertEqual(response.status_code, 200)
        self.client.post(f"/api/documents/{self.doc_id}/open")

    def test_key_is_masked_and_can_be_cleared(self):
        public = self.client.get("/api/config").get_json()
        self.assertTrue(public["has_key"])
        self.assertNotIn("api_key", public)
        self.assertNotIn("sk-abcdefghijklmnop", json.dumps(public))
        self.assertEqual(public["temperature"], 0.0)
        cleared = self.client.post(
            "/api/config",
            json={
                "api_base": "https://api.example.com",
                "api_key": "",
                "clear_key": True,
                "model": "demo",
                "timeout": 20,
                "temperature": 0,
            },
        )
        self.assertEqual(cleared.status_code, 200)
        self.assertFalse(self.client.get("/api/config").get_json()["has_key"])

    def test_chat_exercise_notes_prereq_with_mocked_llm(self):
        with patch("app.llm.chat", return_value="对话回复"):
            chat = self.client.post(
                "/api/chat",
                json={"doc_id": self.doc_id, "question": "这段讲什么"},
            )
        self.assertEqual(chat.status_code, 200)
        self.assertEqual(len(chat.get_json()["messages"]), 2)

        with patch("app.llm.chat", return_value="【题型：非代码题】\n请解释正文。"):
            exercise = self.client.post(
                "/api/exercise", json={"doc_id": self.doc_id, "difficulty": "入门"}
            )
        self.assertEqual(exercise.status_code, 200)
        self.assertIn("请解释", exercise.get_json()["question"])

        with patch("app.llm.chat", return_value="批改结果"):
            review = self.client.post(
                "/api/exercise/submit",
                json={"doc_id": self.doc_id, "code": "我的解答"},
            )
        self.assertEqual(review.status_code, 200)
        self.assertEqual(review.get_json()["review"], "批改结果")
        self.assertEqual(
            self.client.get("/api/session").get_json()["last_submission"]["review"],
            "批改结果",
        )

        with patch("app.llm.chat", return_value='{"suggestions":[{"id":"s1","reason":"补充","text":"要点"}]}'):
            suggest = self.client.post(
                "/api/notes/suggest", json={"doc_id": self.doc_id}
            )
        self.assertEqual(suggest.status_code, 200)
        self.assertEqual(suggest.get_json()["suggestions"][0]["text"], "要点")

        with patch("app.llm.chat", return_value="## 草稿"):
            draft = self.client.post(
                "/api/notes/draft", json={"doc_id": self.doc_id}
            )
        self.assertEqual(draft.status_code, 200)
        self.assertIn("草稿", draft.get_json()["draft"])

        with patch(
            "app.llm.chat",
            return_value='{"focus":"基础","questions":[{"id":"q1","prompt":"会吗？","why":"需要"}]}',
        ):
            prereq = self.client.post("/api/prereq", json={"doc_id": self.doc_id})
        self.assertEqual(prereq.status_code, 200)
        self.assertEqual(len(prereq.get_json()["questions"]), 1)
