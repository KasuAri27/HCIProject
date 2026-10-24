"""外部工具验收：显式选择环境，缺依赖即失败。"""

import argparse
from pathlib import Path
import os
import tempfile
import unittest

from app import converters, parsers


class OcrIntegrationTests(unittest.TestCase):
    def test_real_chinese_and_english_ocr_pipeline(self):
        import pytesseract
        from PIL import Image, ImageDraw, ImageFont
        from reportlab.pdfgen import canvas

        languages = set(pytesseract.get_languages(config=""))
        self.assertTrue({"chi_sim", "eng"}.issubset(languages), languages)
        candidates = [os.environ.get("HCI_CJK_FONT", ""), "C:/Windows/Fonts/simhei.ttf",
                      "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]
        font_path = next((path for path in candidates if path and Path(path).is_file()), None)
        self.assertIsNotNone(font_path, "需要中文字体或 HCI_CJK_FONT")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image = Image.new("RGB", (1800, 500), "white")
            draw = ImageDraw.Draw(image)
            font = ImageFont.truetype(font_path, 90)
            draw.text((70, 100), "学习资料 中文识别", font=font, fill="black")
            draw.text((70, 240), "CI OCR 12345", font=font, fill="black")
            image.save(root / "scan.png")
            pdf = canvas.Canvas(str(root / "scan.pdf"), pagesize=(900, 250))
            pdf.drawImage(str(root / "scan.png"), 0, 0, width=900, height=250)
            pdf.save()
            parsed = parsers.parse_document(root / "scan.pdf", "pdf", "OCR 样本")
            blocks = [block for chapter in parsed["chapters"] for block in chapter["blocks"]]
            text = "".join(block.get("text", "") for block in blocks).replace(" ", "")
            for expected in ("学习资料", "中文识别", "12345"):
                self.assertIn(expected, text)


class OfficeIntegrationTests(unittest.TestCase):
    def test_real_word_and_powerpoint_export_readable_pdf(self):
        from docx import Document
        from pptx import Presentation
        from pypdf import PdfReader

        self.assertTrue(converters.office_available(), "需要 Windows、pywin32 和本机 Office")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            document = Document()
            document.add_paragraph("CI Word 24680 学习资料")
            document.save(root / "sample.docx")
            deck = Presentation()
            slide = deck.slides.add_slide(deck.slide_layouts[1])
            slide.shapes.title.text = "CI PowerPoint 13579"
            slide.placeholders[1].text = "学习资料转换验收"
            deck.save(root / "sample.pptx")
            for suffix, expected in (("docx", "24680"), ("pptx", "13579")):
                with self.subTest(format=suffix):
                    output = converters.convert_to_pdf(root / f"sample.{suffix}", root / f"{suffix}.pdf")
                    reader = PdfReader(output)
                    self.assertGreater(len(reader.pages), 0)
                    self.assertIn(expected, "".join(page.extract_text() or "" for page in reader.pages))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--ocr", action="store_true")
    parser.add_argument("--office", action="store_true")
    args = parser.parse_args()
    if not (args.ocr or args.office):
        parser.error("显式指定 --ocr 或 --office")
    suite = unittest.TestSuite()
    for enabled, case in ((args.ocr, OcrIntegrationTests), (args.office, OfficeIntegrationTests)):
        if enabled:
            suite.addTests(unittest.defaultTestLoader.loadTestsFromTestCase(case))
    return not unittest.TextTestRunner(verbosity=2).run(suite).wasSuccessful()


if __name__ == "__main__":
    raise SystemExit(main())
