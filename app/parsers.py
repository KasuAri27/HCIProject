"""把 Markdown / TXT / DOCX / PPTX / PDF 解析成统一文档结构。

输出结构：
{
  "title": "...",
  "chapters": [
    {"title": "...", "anchor": "...", "source_locator": {...},
     "blocks": [{"type": "heading|paragraph|list|code|quote|table|math|hr",
                 "text": "...", "level": 2, "lang": "...", "rows": [[...]],
                 "source_locator": {...}, "anchor": "..."}]}
  ]
}

内部结构不依赖 PDF 或任何单一格式。每个块带来源位置，阅读与标注只认
block id；回查原文件时使用 source_locator。
"""

import hashlib
import re
from pathlib import Path
from typing import Optional

from . import config, source


class ParseError(Exception):
    """资料解析失败。"""


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _slug(title: str) -> str:
    return source._slugify_heading(title or "") or "section"


def _line_locator(start: int, end: int) -> dict:
    return {"line_start": start, "line_end": end}


def _join_lines(lines: list[str]) -> str:
    return "\n".join(lines).strip()


def _is_fence(line: str) -> Optional[re.Match]:
    return re.match(r"^\s*(```+|~~~+)\s*(.*)$", line)


def _is_table_separator(line: str) -> bool:
    return source._is_table_separator(line)


def _table_rows(lines: list[str]) -> list[list[str]]:
    if len(lines) < 2:
        return []
    header = source._split_cells(lines[0])
    body = [source._split_cells(line) for line in lines[2:]]
    return [header] + body


def _new_chapter(title: str, line_no: int) -> dict:
    return {
        "title": title or "正文",
        "anchor": _slug(title or "正文"),
        "source_locator": _line_locator(line_no, line_no),
        "blocks": [],
    }


def _parse_markdown(path: Path, title: str) -> dict:
    text = _read_text(path).replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")

    # 跳过 YAML frontmatter（保留原始行号）
    start = 0
    if lines and lines[0].strip() == "---":
        for i in range(1, min(len(lines), 200)):
            if lines[i].strip() == "---":
                start = i + 1
                break

    chapters: list[dict] = []
    current: Optional[dict] = None
    anchors: dict[str, int] = {}

    def ensure_chapter(chapter_title: str, line_no: int) -> dict:
        nonlocal current
        if current is None:
            current = _new_chapter(chapter_title, line_no)
            chapters.append(current)
        return current

    def add_block(block: dict) -> None:
        nonlocal current
        if current is None:
            current = _new_chapter(title or "正文", block.get("source_locator", {}).get("line_start", 1))
            chapters.append(current)
        current["blocks"].append(block)

    def anchor_for(raw: str) -> str:
        base = _slug(raw)
        n = anchors.get(base, 0)
        anchors[base] = n + 1
        return base if n == 0 else f"{base}-{n}"

    i = start
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue

        fence = _is_fence(line)
        if fence:
            marker = fence.group(1)
            lang = (fence.group(2) or "").strip()
            j = i + 1
            body: list[str] = []
            while j < len(lines) and not re.match(r"^\s*" + re.escape(marker[0]) + r"{3,}\s*$", lines[j]):
                body.append(lines[j])
                j += 1
            end = min(j, len(lines) - 1) + 1
            ensure_chapter("代码" if current is None else current["title"], i + 1)
            add_block(
                {
                    "type": "code",
                    "text": "\n".join(body),
                    "lang": lang,
                    "source_locator": _line_locator(i + 1, end),
                }
            )
            i = j + 1 if j < len(lines) else j
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            raw_title = heading.group(2).strip()
            block = {
                "type": "heading",
                "text": raw_title,
                "level": level,
                "anchor": anchor_for(raw_title),
                "source_locator": _line_locator(i + 1, i + 1),
            }
            if level <= 2:
                current = _new_chapter(raw_title, i + 1)
                chapters.append(current)
            else:
                ensure_chapter(title or "正文", i + 1)
            add_block(block)
            i += 1
            continue

        if line.strip() in ("$$", "\\["):
            j = i + 1
            body = []
            while j < len(lines) and lines[j].strip() not in ("$$", "\\]"):
                body.append(lines[j])
                j += 1
            end = min(j, len(lines) - 1) + 1
            add_block(
                {
                    "type": "math",
                    "text": "\n".join(body).strip(),
                    "source_locator": _line_locator(i + 1, end),
                }
            )
            i = j + 1 if j < len(lines) else j
            continue

        display_math = re.match(r"^\s*\$\$(.+)\$\$\s*$", line)
        if display_math:
            add_block(
                {
                    "type": "math",
                    "text": display_math.group(1).strip(),
                    "source_locator": _line_locator(i + 1, i + 1),
                }
            )
            i += 1
            continue

        if re.match(r"^\s*([-*_])\s*\1\s*\1[\s\S]*$", line) or re.match(r"^\s*---+\s*$", line):
            add_block({"type": "hr", "text": "", "source_locator": _line_locator(i + 1, i + 1)})
            i += 1
            continue

        if line.lstrip().startswith("|") and i + 1 < len(lines) and _is_table_separator(lines[i + 1]):
            j = i + 1
            while j + 1 < len(lines) and lines[j + 1].lstrip().startswith("|"):
                j += 1
            table_lines = lines[i : j + 1]
            add_block(
                {
                    "type": "table",
                    "text": _join_lines(table_lines),
                    "rows": _table_rows(table_lines),
                    "source_locator": _line_locator(i + 1, j + 1),
                }
            )
            i = j + 1
            continue

        if re.match(r"^\s*>\s?", line):
            j = i
            body: list[str] = []
            while j < len(lines) and re.match(r"^\s*>\s?", lines[j]):
                body.append(re.sub(r"^\s*>\s?", "", lines[j]))
                j += 1
            add_block(
                {
                    "type": "quote",
                    "text": _join_lines(body),
                    "source_locator": _line_locator(i + 1, j),
                }
            )
            i = j
            continue

        if re.match(r"^\s*([-*+]|\d+\.)\s+", line):
            j = i
            body: list[str] = []
            while j < len(lines) and re.match(r"^\s*([-*+]|\d+\.)\s+", lines[j]):
                body.append(lines[j])
                j += 1
            add_block(
                {
                    "type": "list",
                    "text": _join_lines(body),
                    "source_locator": _line_locator(i + 1, j),
                }
            )
            i = j
            continue

        # 段落：收集到空行或下一个块起始
        j = i + 1
        body = [line]
        while j < len(lines):
            nxt = lines[j]
            if not nxt.strip():
                break
            if (
                _is_fence(nxt)
                or re.match(r"^(#{1,6})\s+", nxt)
                or nxt.strip() in ("$$", "\\[")
                or re.match(r"^\s*\$\$(.+)\$\$\s*$", nxt)
                or re.match(r"^\s*>\s?", nxt)
                or re.match(r"^\s*([-*+]|\d+\.)\s+", nxt)
                or (nxt.lstrip().startswith("|") and j + 1 < len(lines) and _is_table_separator(lines[j + 1]))
            ):
                break
            body.append(nxt)
            j += 1
        add_block(
            {
                "type": "paragraph",
                "text": _join_lines(body),
                "source_locator": _line_locator(i + 1, j),
            }
        )
        i = j

    if not chapters:
        chapters = [_new_chapter(title or "正文", 1)]
    return {"title": title or path.stem, "chapters": chapters}


def _parse_txt(path: Path, title: str) -> dict:
    text = _read_text(path).replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    chapter = _new_chapter(title or path.stem, 1)
    i = 0
    while i < len(lines):
        if not lines[i].strip():
            i += 1
            continue
        j = i + 1
        while j < len(lines) and lines[j].strip():
            j += 1
        chapter["blocks"].append(
            {
                "type": "paragraph",
                "text": _join_lines(lines[i:j]),
                "source_locator": _line_locator(i + 1, j),
            }
        )
        i = j
    if not chapter["blocks"]:
        chapter["blocks"].append(
            {"type": "paragraph", "text": "", "source_locator": _line_locator(1, 1)}
        )
    return {"title": title or path.stem, "chapters": [chapter]}


def _docx_heading_level(style_name: str) -> Optional[int]:
    m = re.search(r"(\d+)$", style_name or "")
    if m and ("heading" in (style_name or "").lower() or "标题" in (style_name or "")):
        return max(1, min(6, int(m.group(1))))
    return None


def _parse_docx(path: Path, title: str) -> dict:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - dependency error path
        raise ParseError("缺少 python-docx，无法解析 DOCX") from exc

    try:
        document = Document(str(path))
    except Exception as exc:
        raise ParseError(f"DOCX 打开失败：{exc}") from exc

    chapters: list[dict] = []
    current: Optional[dict] = None
    anchors: dict[str, int] = {}
    paragraph_no = 0
    table_no = 0

    def anchor_for(raw: str) -> str:
        base = _slug(raw)
        n = anchors.get(base, 0)
        anchors[base] = n + 1
        return base if n == 0 else f"{base}-{n}"

    def ensure_chapter(chapter_title: str, locator: dict) -> dict:
        nonlocal current
        if current is None:
            current = {
                "title": chapter_title or title or path.stem,
                "anchor": _slug(chapter_title or title or path.stem),
                "source_locator": locator,
                "blocks": [],
            }
            chapters.append(current)
        return current

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            paragraph_no += 1
            p = Paragraph(child, document)
            text = (p.text or "").strip()
            style = (p.style.name if p.style is not None else "") or ""
            if not text:
                # 图片等无文字内容：保留一个可见提示，避免静默丢内容
                try:
                    has_image = bool(child.xpath(".//a:blip"))
                except Exception:
                    has_image = False
                if has_image:
                    ensure_chapter(current["title"] if current else title, {"paragraph": paragraph_no})
                    current["blocks"].append(
                        {
                            "type": "paragraph",
                            "text": "[图片]",
                            "source_locator": {"paragraph": paragraph_no},
                        }
                    )
                continue
            level = _docx_heading_level(style)
            if level:
                block = {
                    "type": "heading",
                    "text": text,
                    "level": level,
                    "anchor": anchor_for(text),
                    "source_locator": {"paragraph": paragraph_no},
                }
                if level <= 2:
                    current = {
                        "title": text,
                        "anchor": anchor_for(text),
                        "source_locator": {"paragraph": paragraph_no},
                        "blocks": [],
                    }
                    chapters.append(current)
                else:
                    ensure_chapter(title, {"paragraph": paragraph_no})
                current["blocks"].append(block)
                continue
            ensure_chapter(title, {"paragraph": paragraph_no})
            is_list = "list" in style.lower() or "列表" in style
            current["blocks"].append(
                {
                    "type": "list" if is_list else "paragraph",
                    "text": text,
                    "source_locator": {"paragraph": paragraph_no},
                }
            )
        elif tag == "tbl":
            table_no += 1
            table = Table(child, document)
            rows = []
            for row in table.rows:
                rows.append([(cell.text or "").strip() for cell in row.cells])
            ensure_chapter(current["title"] if current else title, {"table": table_no})
            current["blocks"].append(
                {
                    "type": "table",
                    "text": "\n".join(" | ".join(row) for row in rows),
                    "rows": rows,
                    "source_locator": {"table": table_no},
                }
            )

    if not chapters:
        chapters = [
            {
                "title": title or path.stem,
                "anchor": _slug(title or path.stem),
                "source_locator": {},
                "blocks": [],
            }
        ]
    return {"title": title or path.stem, "chapters": chapters}


def _pptx_save_image(blob: bytes, ext: str, asset_dir, filename: str, seen: dict) -> str:
    """保存幻灯片图片，返回文件名；同一图片只写一次。"""
    if not asset_dir:
        return ""
    digest = hashlib.sha256(blob).hexdigest()
    if digest in seen:
        return seen[digest]
    safe_ext = re.sub(r"[^a-z0-9]", "", (ext or "png").lower()) or "png"
    name = f"{filename}.{safe_ext}"
    folder = Path(asset_dir)
    folder.mkdir(parents=True, exist_ok=True)
    (folder / name).write_bytes(blob)
    seen[digest] = name
    return name


def _pptx_shapes(shapes):
    """展平组合形状。"""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        shape_type = getattr(shape, "shape_type", None)
        if shape_type is not None and shape_type == MSO_SHAPE_TYPE.GROUP and hasattr(shape, "shapes"):
            yield from _pptx_shapes(shape.shapes)
        else:
            yield shape

def _parse_pptx(path: Path, title: str, asset_dir=None, asset_url_prefix: str = "") -> dict:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as exc:  # pragma: no cover - dependency error path
        raise ParseError("缺少 python-pptx，无法解析 PPTX") from exc

    try:
        deck = Presentation(str(path))
    except Exception as exc:
        raise ParseError(f"PPTX 打开失败：{exc}") from exc

    chapters: list[dict] = []
    seen_images: dict[str, str] = {}

    for slide_no, slide in enumerate(deck.slides, start=1):
        title_shape = slide.shapes.title
        slide_title = ""
        if title_shape is not None:
            slide_title = (title_shape.text or "").strip()
        chapter_title = slide_title or f"第 {slide_no} 页"
        chapter = {
            "title": chapter_title,
            "anchor": f"slide-{slide_no}",
            "source_locator": {"slide": slide_no},
            "blocks": [],
        }
        if slide_title:
            chapter["blocks"].append(
                {
                    "type": "heading",
                    "text": slide_title,
                    "level": 2,
                    "source_locator": {"slide": slide_no},
                }
            )

        title_id = getattr(title_shape, "shape_id", None) if title_shape is not None else None
        for shape_no, shape in enumerate(_pptx_shapes(slide.shapes), start=1):
            if title_id is not None and getattr(shape, "shape_id", None) == title_id:
                continue
            locator = {"slide": slide_no, "shape": shape_no}

            if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = shape.image
                    name = _pptx_save_image(
                        image.blob,
                        image.ext,
                        asset_dir,
                        f"slide{slide_no}_{shape_no}",
                        seen_images,
                    )
                except Exception as exc:
                    chapter["blocks"].append(
                        {
                            "type": "paragraph",
                            "text": f"[图片读取失败：{exc}]",
                            "source_locator": locator,
                        }
                    )
                    continue
                if name:
                    url = f"{asset_url_prefix}{name}" if asset_url_prefix else name
                    alt = (getattr(shape, "name", "") or "幻灯片图片").strip()
                    alt = alt.replace("[", " ").replace("]", " ").strip() or "幻灯片图片"
                    chapter["blocks"].append(
                        {
                            "type": "image",
                            "text": f"[图片] {alt}",
                            "markdown": f"![{alt}]({url})",
                            "source_locator": locator,
                        }
                    )
                else:
                    chapter["blocks"].append(
                        {
                            "type": "paragraph",
                            "text": "[图片] 未能保存图片文件",
                            "source_locator": locator,
                        }
                    )
                continue

            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    rows.append([(cell.text or "").strip() for cell in row.cells])
                if rows:
                    chapter["blocks"].append(
                        {
                            "type": "table",
                            "text": "\n".join(" | ".join(row) for row in rows),
                            "rows": rows,
                            "source_locator": locator,
                        }
                    )
                continue

            if getattr(shape, "has_chart", False):
                chart_title = ""
                try:
                    chart_title = (shape.chart.chart_title.text_frame.text or "").strip()
                except Exception:
                    chart_title = ""
                label = f"图表：{chart_title}" if chart_title else "图表"
                chapter["blocks"].append(
                    {
                        "type": "paragraph",
                        "text": f"[{label}] 图表数据未解析，请在原文件中查看。",
                        "source_locator": locator,
                    }
                )
                continue

            if getattr(shape, "has_text_frame", False):
                paragraphs = []
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs).strip()
                    if not text:
                        text = (paragraph.text or "").strip()
                    if not text:
                        continue
                    level = min(3, max(0, getattr(paragraph, "level", 0) or 0))
                    paragraphs.append(("  " * level) + text)
                if not paragraphs:
                    continue
                block_type = "paragraph" if len(paragraphs) == 1 else "list"
                chapter["blocks"].append(
                    {
                        "type": block_type,
                        "text": "\n".join(paragraphs),
                        "source_locator": locator,
                    }
                )

        try:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
        except Exception:
            notes = ""
        if notes:
            chapter["blocks"].append(
                {
                    "type": "heading",
                    "text": "备注",
                    "level": 3,
                    "source_locator": {"slide": slide_no, "note": True},
                }
            )
            chapter["blocks"].append(
                {
                    "type": "paragraph",
                    "text": notes,
                    "source_locator": {"slide": slide_no, "note": True},
                }
            )

        if not chapter["blocks"]:
            chapter["blocks"].append(
                {
                    "type": "paragraph",
                    "text": "（本页没有可提取的文字、表格或图片）",
                    "source_locator": {"slide": slide_no},
                }
            )
        chapters.append(chapter)

    if not chapters:
        raise ParseError("PPTX 没有幻灯片")
    return {"title": title or path.stem, "chapters": chapters}


def _pdf_paragraphs(text: str) -> list[str]:
    text = text.replace("\xa0", " ").replace("\u3000", " ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    parts = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if len(parts) <= 1:
        lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
        parts = lines if len(lines) > 1 else parts
    return parts


def _ocr_pdf_page(page) -> tuple[str, str]:
    """对扫描页的嵌入图片做 OCR，返回 (文本, 错误说明)。"""
    try:
        from io import BytesIO

        import pytesseract
        from PIL import Image
    except ImportError:
        return "", "OCR 依赖未安装（需要 pytesseract 与 Pillow）"

    try:
        images = [image.data for image in page.images]
    except Exception as exc:
        return "", f"读取页面图片失败：{exc}"
    if not images:
        return "", "页面没有可识别的嵌入图片"
    data = max(images, key=len)
    try:
        img = Image.open(BytesIO(data))
        img.load()
        if img.mode not in ("RGB", "L"):
            img = img.convert("RGB")
        text = pytesseract.image_to_string(img, lang=config.OCR_LANG)
    except Exception as exc:
        message = str(exc)
        lower = message.lower()
        if "tesseract is not installed" in lower or "not in your path" in lower:
            return "", "未找到 Tesseract OCR 引擎，请安装后加入 PATH"
        if "failed loading language" in lower or "couldn't load any languages" in lower:
            return "", f"Tesseract 缺少语言包：{config.OCR_LANG}"
        return "", f"OCR 失败：{message[:120]}"
    return (text or "").strip(), ""


def _parse_pdf(path: Path, title: str) -> dict:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ParseError("缺少 pypdf，无法解析 PDF") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise ParseError(f"PDF 打开失败：{exc}") from exc

    chapters: list[dict] = []
    has_text = False
    ocr_errors: list[str] = []
    for page_no, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            raise ParseError(f"第 {page_no} 页文本提取失败：{exc}") from exc
        paragraphs = _pdf_paragraphs(text)
        page_ocr = False
        if len((text or "").strip()) < 20:
            ocr_text, ocr_error = _ocr_pdf_page(page)
            if ocr_text:
                paragraphs = _pdf_paragraphs(ocr_text)
                page_ocr = True
            elif ocr_error:
                ocr_errors.append(f"第 {page_no} 页：{ocr_error}")
        chapter = {
            "title": f"第 {page_no} 页",
            "anchor": f"page-{page_no}",
            "source_locator": {"page": page_no},
            "blocks": [],
        }
        if paragraphs:
            has_text = True
            for idx, paragraph in enumerate(paragraphs):
                locator = {"page": page_no, "paragraph": idx + 1}
                if page_ocr:
                    locator["ocr"] = True
                chapter["blocks"].append(
                    {
                        "type": "paragraph",
                        "text": paragraph,
                        "source_locator": locator,
                    }
                )
        if not paragraphs or (page_ocr is False and len((text or "").strip()) < 20):
            reason = next(
                (e.split("：", 1)[1] for e in ocr_errors if e.startswith(f"第 {page_no} 页：")),
                "",
            )
            hint = "（本页为扫描页，未识别出文字）"
            if reason:
                hint = f"（本页为扫描页，OCR 未成功：{reason}）"
            chapter["blocks"].append(
                {
                    "type": "paragraph",
                    "text": hint,
                    "source_locator": {"page": page_no},
                }
            )
        chapters.append(chapter)

    if not has_text:
        detail = ""
        if ocr_errors:
            detail = f"（{ocr_errors[0]}）"
        raise ParseError(
            "PDF 没有可提取文本；扫描页 OCR 未成功"
            + detail
            + "。如为扫描件，请安装 Tesseract 与 chi_sim 语言包后重试"
        )
    if not chapters:
        raise ParseError("PDF 没有页面")
    return {"title": title or path.stem, "chapters": chapters}



_MD_ESCAPE_RE = re.compile(r"([\\`*_{}\[\]#+\-.!>|$~])")


def _escape_md(text: str) -> str:
    """把纯文本转成 Markdown 时不触发格式语法。"""
    return _MD_ESCAPE_RE.sub(lambda m: chr(92) + m.group(1), text or "")


def _escape_md_lines(text: str) -> str:
    return "\n".join(_escape_md(line) for line in (text or "").splitlines())


def _markdown_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)

    def cell(value) -> str:
        return _escape_md(str(value).replace("|", "\\|").replace("\n", " ").strip())

    def padded(row: list[str]) -> str:
        values = [cell(v) for v in row] + [""] * (width - len(row))
        return "| " + " | ".join(values) + " |"

    lines = [padded(rows[0]), "| " + " | ".join("---" for _ in range(width)) + " |"]
    lines.extend(padded(row) for row in rows[1:])
    return "\n".join(lines)


def attach_markdown(parsed: dict, source_format: str) -> dict:
    """给每个块补上规范化 Markdown，并生成整篇 Markdown。"""
    is_markdown = source_format in ("md", "markdown")
    for chapter in parsed.get("chapters") or []:
        for block in chapter.get("blocks") or []:
            btype = block.get("type") or "paragraph"
            text = block.get("text") or ""
            if btype == "image":
                block["markdown"] = block.get("markdown") or f"![图片]({text})"
                continue
            if is_markdown:
                if btype == "heading":
                    level = max(1, min(6, int(block.get("level") or 2)))
                    md = "#" * level + " " + text
                elif btype == "code":
                    lang = (block.get("lang") or "").strip()
                    md = f"```{lang}\n{text}\n```"
                elif btype == "math":
                    md = f"$$\n{text}\n$$"
                elif btype == "quote":
                    md = chr(10).join("> " + line for line in text.splitlines())
                elif btype == "hr":
                    md = "---"
                else:
                    md = text
            else:
                if btype == "heading":
                    level = max(1, min(6, int(block.get("level") or 2)))
                    md = "#" * level + " " + _escape_md(" ".join(text.splitlines()))
                elif btype == "list":
                    md = "\n".join(
                        "- " + _escape_md(line) for line in text.splitlines() if line.strip()
                    )
                elif btype == "quote":
                    md = chr(10).join("> " + line for line in text.splitlines())
                elif btype == "hr":
                    md = "---"
                else:
                    md = _escape_md_lines(text)
            block["markdown"] = md
    parsed["markdown"] = "\n\n".join(
        block.get("markdown") or ""
        for chapter in parsed.get("chapters") or []
        for block in chapter.get("blocks") or []
    ).strip()
    return parsed


def parse_document(
    path: Path,
    source_format: str,
    title: str = "",
    asset_dir=None,
    asset_url_prefix: str = "",
) -> dict:
    """按扩展名解析一份资料，返回统一结构。"""
    fmt = (source_format or path.suffix.lstrip(".")).lower()
    if fmt in ("md", "markdown"):
        parsed = _parse_markdown(path, title)
    elif fmt == "txt":
        parsed = _parse_txt(path, title)
    elif fmt == "docx":
        parsed = _parse_docx(path, title)
    elif fmt == "pptx":
        parsed = _parse_pptx(path, title, asset_dir=asset_dir, asset_url_prefix=asset_url_prefix)
    elif fmt == "pdf":
        parsed = _parse_pdf(path, title)
    else:
        raise ParseError(f"不支持的文件类型：{fmt}")
    return attach_markdown(parsed, fmt)


def locator_label(source_format: str, locator: dict) -> str:
    """把来源位置转成用户可读的定位说明。"""
    locator = locator or {}
    fmt = (source_format or "").lower()
    if "slide" in locator:
        label = f"第 {locator['slide']} 页"
        if locator.get("note"):
            label += " · 备注"
        elif locator.get("shape"):
            label += f" · 形状 {locator['shape']}"
        return label
    if "page" in locator:
        label = f"第 {locator['page']} 页"
        if locator.get("paragraph"):
            label += f" · 第 {locator['paragraph']} 段"
        if locator.get("ocr"):
            label += " · OCR"
        return label
    if "line_start" in locator:
        start = locator["line_start"]
        end = locator.get("line_end") or start
        return f"第 {start} 行" if start == end else f"第 {start}–{end} 行"
    if "paragraph" in locator:
        return f"第 {locator['paragraph']} 段"
    if "table" in locator:
        return f"表格 {locator['table']}"
    if fmt in ("md", "markdown"):
        return "Markdown"
    if fmt == "txt":
        return "文本"
    return ""
